"""T1 situations and assessment versions from real NBS measurements (AS-011), on PostgreSQL."""

from __future__ import annotations

from collections.abc import Iterator
from datetime import UTC, date, datetime

import openpyxl
import pytest
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from africasignal.models import (
    AssessmentInput,
    AssessmentVersion,
    Measurement,
    Place,
    Situation,
    Source,
)
from africasignal.publish.situations import (
    POLICY_UNAPPLIED,
    assess_situation,
    ensure_situations,
    request_assessments,
    slug_for,
    source_short_name,
)
from africasignal.storage import S3Store
from tests.integration.nbs_support import add_places, add_source, import_bytes, make_store
from tests.unit.sources.nbs_fixtures import FIXTURES, edit

PMS_SEP, PMS_OCT = "FUEL_SEPT_2024_REPORT.xlsx", "PMS_OCT_2024_REPORT.xlsx"
WHEN = datetime(2024, 11, 25, 12, tzinfo=UTC)  # a week after the October 2024 petrol release
TODAY = datetime(2026, 9, 30, 12, tzinfo=UTC)


@pytest.fixture
def store() -> Iterator[S3Store]:
    yield from make_store()


@pytest.fixture
def places(session: Session) -> dict[str, int]:
    return add_places(session)


@pytest.fixture
def source(session: Session, places: dict[str, int]) -> Source:
    return add_source(session)


def _situation(session: Session, item: str, place_code: str) -> Situation:
    return session.scalars(
        select(Situation).where(Situation.slug == slug_for(item, place_code))
    ).one()


def _versions(session: Session, situation: Situation) -> list[AssessmentVersion]:
    return list(
        session.scalars(
            select(AssessmentVersion)
            .where(AssessmentVersion.situation_id == situation.id)
            .order_by(AssessmentVersion.version)
        )
    )


def _all_pairs(session: Session, item: str) -> set[tuple[str, int]]:
    return {(item, p) for p in session.scalars(select(Place.id).where(Place.kind != "lga"))}


def test_source_short_name_is_the_acronym_of_the_owner(source: Source) -> None:
    source.owner = "National Bureau of Statistics"
    assert source_short_name(source) == "NBS"
    source.owner = "Nigerian Midstream and Downstream Petroleum Regulatory Authority"
    assert source_short_name(source) == "NMDPRA"
    source.owner = "Punch"
    assert source_short_name(source) == "Punch"


# --- situations (B8.1) -------------------------------------------------------------------------


def test_one_situation_per_item_and_place_with_two_consecutive_months(
    session: Session, store: S3Store, source: Source, places: dict[str, int]
) -> None:
    import_bytes(session, store, source, PMS_OCT)
    found = ensure_situations(session, _all_pairs(session, "pms_litre"))
    assert len(found) == 38  # the country and 37 states
    lagos = _situation(session, "pms_litre", "NG-LA")
    assert lagos.slug == "price-pms_litre-ng-la"  # the slug the spec gives as its example
    assert (lagos.kind, lagos.topic, lagos.status) == ("price_series", "energy", "active")
    assert lagos.title == "Petrol (PMS) price in Lagos State"
    assert lagos.place_id == places["NG-LA"] and lagos.item_code == "pms_litre"
    assert _situation(session, "pms_litre", "NG").title == "Petrol (PMS) price in Nigeria"
    assert _situation(session, "pms_litre", "NG-FC").title == (
        "Petrol (PMS) price in the Federal Capital Territory"
    )
    assert lagos.current_version_id is None  # nothing is published before the policy (AS-012)


def test_creating_situations_twice_changes_nothing(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_OCT)
    pairs = _all_pairs(session, "pms_litre")
    first = {s.id for s in ensure_situations(session, pairs)}
    second = {s.id for s in ensure_situations(session, pairs)}
    assert first == second
    assert session.scalar(select(func.count()).select_from(Situation)) == 38


def test_food_items_only_have_a_national_situation(
    session: Session, store: S3Store, source: Source, places: dict[str, int]
) -> None:
    import_bytes(session, store, source, "selected_food_oct_2024.xlsx")
    found = ensure_situations(session, _all_pairs(session, "rice_local_1kg"))
    assert [s.slug for s in found] == ["price-rice_local_1kg-ng"]


def test_a_place_with_a_single_month_gets_no_situation(
    session: Session, store: S3Store, source: Source, places: dict[str, int]
) -> None:
    import_bytes(session, store, source, PMS_OCT)
    lagos = places["NG-LA"]
    # Keep only October 2024 for Lagos: no consecutive pair is left.
    session.execute(
        text("DELETE FROM measurement WHERE place_id = :p AND period_start <> '2024-10-01'"),
        {"p": lagos},
    )
    found = {s.slug for s in ensure_situations(session, [("pms_litre", lagos)])}
    assert found == set()
    assert ensure_situations(session, [("pms_litre", places["NG-OG"])])  # Ogun still qualifies


def test_unknown_items_and_non_scope_places_are_ignored(
    session: Session, store: S3Store, source: Source, places: dict[str, int]
) -> None:
    import_bytes(session, store, source, PMS_OCT)
    lga = Place(kind="lga", name="Ikeja", code="NG-LA-IKE", parent_id=places["NG-LA"])
    session.add(lga)
    session.flush()
    assert (
        ensure_situations(session, [("no_such_item", places["NG-LA"]), ("pms_litre", lga.id)]) == []
    )


# --- assessments -------------------------------------------------------------------------------


def test_a_state_assessment_from_real_petrol_data(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_SEP)
    import_bytes(session, store, source, PMS_OCT)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    outcome = assess_situation(session, lagos.id, WHEN)
    assert outcome.outcome == "created"
    v = outcome.version
    assert v is not None
    assert v.headline == (
        "Average petrol (PMS) price in Lagos State rose 8.0% in October 2024 to ₦1,080.95 (NBS)"
    )
    assert (v.version, v.status, v.policy_version) == (1, "draft", POLICY_UNAPPLIED)
    assert (v.template, v.template_version) == ("T1_price_change", "T1-1")
    # +82.9 % on the year is 4.1x the 20 % threshold (above 4x), so severity is high
    assert (v.evidence_state, v.severity) == ("reported", "high")
    assert v.scope_label == "Lagos State (state average, NBS)" and v.period_label == "October 2024"
    assert v.valid_until == datetime(2025, 1, 14, 23, 59, 59, tzinfo=UTC)
    assert v.last_checked_at == WHEN and v.explanation is None and v.published_at is None
    assert v.supersedes_id is None and v.change_summary is None and v.withheld_reasons == []
    facts = {f["label"]: f["value"] for f in v.facts}
    assert facts["Current price"] == 1080.95 and facts["Previous month"] == 1000.48
    assert facts["Same month last year"] == 590.95
    assert [f["factor"] for f in v.possible_factors][:2] == ["Crude oil price", "Exchange rate"]
    assert all(f["status"] == "not_checked" for f in v.possible_factors)
    session.refresh(lagos)
    assert lagos.current_version_id is None


def test_the_national_assessment_matches_what_nbs_states_in_its_summary(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_SEP)
    import_bytes(session, store, source, PMS_OCT)
    (national,) = ensure_situations(session, [("pms_litre", _place(session, "NG"))])
    v = assess_situation(session, national.id, WHEN).version
    assert v is not None
    # NBS: "increased by 14.98 %" on the month and "87.88 %" on the year: 15.0 and 87.9 rounded.
    assert v.headline == (
        "Average petrol (PMS) price in Nigeria rose 15.0% in October 2024 to ₦1,184.83 (NBS)"
    )
    facts = {f["label"]: f["value"] for f in v.facts}
    assert facts["Year-on-year change"] == 87.9
    # NBS: highest Ebonyi 1292.86, Jigawa 1288.18, Borno 1283.79; lowest Delta 1050.00,
    # Nasarawa 1063.68, Lagos 1080.95.
    assert facts["Highest states"] == [
        {"place": "Ebonyi", "value": 1292.86},
        {"place": "Jigawa", "value": 1288.18},
        {"place": "Borno", "value": 1283.79},
    ]
    assert facts["Lowest states"] == [
        {"place": "Delta", "value": 1050.0},
        {"place": "Nasarawa", "value": 1063.68},
        {"place": "Lagos", "value": 1080.95},
    ]
    assert facts["States up"] + facts["States down"] + facts["States unchanged"] == 37


def test_states_up_and_down_agree_with_an_independent_reading_of_the_raw_cells(
    session: Session, store: S3Store, source: Source
) -> None:
    """Counts the states from the workbook cells directly, without the parser."""
    import_bytes(session, store, source, PMS_OCT)
    ws = openpyxl.load_workbook(FIXTURES / PMS_OCT, data_only=True)["Fuel_October 2024"]
    up = down = flat = 0
    for row in range(16, 53):  # the 37 state rows; column C is September, D is October
        change = (ws.cell(row, 4).value / ws.cell(row, 3).value - 1) * 100
        if abs(change) <= 0.5:
            flat += 1
        elif change > 0:
            up += 1
        else:
            down += 1
    assert (up, down, flat) == (36, 0, 1)  # only Katsina (1096.15 -> 1100.00) is within 0.5 %
    (national,) = ensure_situations(session, [("pms_litre", _place(session, "NG"))])
    v = assess_situation(session, national.id, WHEN).version
    assert v is not None
    facts = {f["label"]: f["value"] for f in v.facts}
    assert (facts["States up"], facts["States down"], facts["States unchanged"]) == (up, down, flat)


def test_reassessing_with_the_same_inputs_stores_nothing_new(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_OCT)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    first = assess_situation(session, lagos.id, WHEN).version
    again = assess_situation(session, lagos.id, TODAY)
    assert first is not None and again.outcome == "unchanged" and again.version is first
    assert len(_versions(session, lagos)) == 1
    assert first.last_checked_at == TODAY  # checked again; nothing changed
    assert first.valid_until == datetime(2025, 1, 14, 23, 59, 59, tzinfo=UTC)  # not extended


def test_same_hash_even_though_the_clock_moved_into_staleness(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_OCT)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    fresh = assess_situation(session, lagos.id, WHEN).version
    assert fresh is not None and fresh.evidence_state == "reported"
    assess_situation(session, lagos.id, TODAY)
    assert fresh.evidence_state == "reported"  # expiry (AS-013), not re-assessment, makes it stale
    assert len(_versions(session, lagos)) == 1


def test_data_that_is_already_stale_when_first_assessed_is_insufficient(
    session: Session, store: S3Store, source: Source
) -> None:
    """NBS's newest price watch is October 2024; assessing it on 30 September 2026."""
    import_bytes(session, store, source, PMS_OCT)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    v = assess_situation(session, lagos.id, TODAY).version
    assert v is not None
    assert (v.evidence_state, v.severity) == ("insufficient", "none")
    assert any("more than 120 days ago" in u for u in v.unknowns)


def test_a_new_release_creates_version_two_linked_to_version_one(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_SEP)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    v1 = assess_situation(session, lagos.id, datetime(2024, 10, 20, tzinfo=UTC)).version
    assert v1 is not None and v1.period_label == "September 2024"
    import_bytes(session, store, source, PMS_OCT)
    v2 = assess_situation(session, lagos.id, WHEN).version
    assert v2 is not None and v2.id != v1.id
    assert (v2.version, v2.supersedes_id) == (2, v1.id)
    assert v2.change_summary == "Updated with October 2024 data (previously September 2024)"
    assert v1.status == "draft"  # superseding is the publication step's job (AS-012/013)
    assert v1.inputs_hash != v2.inputs_hash


def test_a_restated_value_creates_a_new_version_because_the_inputs_changed(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_SEP)
    import_bytes(session, store, source, PMS_OCT)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    v1 = assess_situation(session, lagos.id, WHEN).version
    restated = edit(PMS_OCT, "Fuel_October 2024", {"C40": 1010.48})  # Lagos' September value
    import_bytes(session, store, source, PMS_OCT, restated)
    v2 = assess_situation(session, lagos.id, WHEN).version
    assert v1 is not None and v2 is not None and v2.version == 2
    assert v2.headline.endswith("rose 7.0% in October 2024 to ₦1,080.95 (NBS)")  # 1080.95/1010.48
    assert v2.change_summary == "Re-assessed after the inputs changed"


def test_versions_record_every_input_for_invalidation_lookups(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_OCT)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    v = assess_situation(session, lagos.id, WHEN).version
    assert v is not None
    inputs = session.execute(
        select(AssessmentInput.input_kind, AssessmentInput.input_id).where(
            AssessmentInput.assessment_version_id == v.id
        )
    ).all()
    measurements = {i for k, i in inputs if k == "measurement"}
    docs = {i for k, i in inputs if k == "evidence_document"}
    current = session.scalars(
        select(Measurement).where(
            Measurement.place_id == lagos.place_id, Measurement.superseded_by_id.is_(None)
        )
    ).all()
    wanted = {
        m.id
        for m in current
        if m.period_start in (date(2024, 10, 1), date(2024, 9, 1), date(2023, 10, 1))
    }
    assert measurements == wanted and len(docs) == 1
    # the index the invalidation step will query
    hit = session.scalars(
        select(AssessmentInput.assessment_version_id).where(
            AssessmentInput.input_kind == "measurement",
            AssessmentInput.input_id == min(measurements),
        )
    ).all()
    assert hit == [v.id]


def test_measurements_from_a_withdrawn_document_are_not_used(
    session: Session, store: S3Store, source: Source
) -> None:
    doc, _ = import_bytes(session, store, source, PMS_OCT)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    doc.status = "withdrawn"
    session.flush()
    outcome = assess_situation(session, lagos.id, WHEN)
    assert (outcome.outcome, outcome.reason) == ("skipped", "no current measurements")
    assert _versions(session, lagos) == []


def test_assessing_a_situation_that_does_not_exist_is_a_quiet_skip(session: Session) -> None:
    assert assess_situation(session, 424242, WHEN).outcome == "skipped"


def test_one_release_is_enough_because_it_carries_the_previous_month(
    session: Session, store: S3Store, source: Source
) -> None:
    import_bytes(session, store, source, PMS_SEP)
    (lagos,) = ensure_situations(session, [("pms_litre", _place(session, "NG-LA"))])
    v = assess_situation(session, lagos.id, datetime(2024, 10, 20, tzinfo=UTC)).version
    assert v is not None and v.period_label == "September 2024"
    facts = {f["label"]: f["value"] for f in v.facts}
    assert (facts["Current price"], facts["Previous month"]) == (1000.48, 765.29)  # Sep and Aug
    assert v.headline.endswith("rose 30.7% in September 2024 to ₦1,000.48 (NBS)")


# --- queueing ----------------------------------------------------------------------------------


def _assess_jobs(session: Session) -> list[str]:
    return list(
        session.execute(
            text("SELECT dedupe_key FROM job WHERE kind = 'assess_situation' ORDER BY id")
        ).scalars()
    )


def test_changed_values_queue_one_assessment_per_situation_and_document(
    session: Session, store: S3Store, source: Source
) -> None:
    doc, result = import_bytes(session, store, source, PMS_OCT)
    queued = request_assessments(session, result.touched, doc.id)
    assert len(queued) == 38  # the country and 37 states
    keys = _assess_jobs(session)
    assert len(keys) == 38 and all(k.endswith(f":{doc.id}") for k in keys)
    assert request_assessments(session, result.touched, doc.id) == []  # same document: nothing new
    assert len(_assess_jobs(session)) == 38


def test_a_new_document_queues_its_own_assessments(
    session: Session, store: S3Store, source: Source
) -> None:
    sep_doc, sep = import_bytes(session, store, source, PMS_SEP)
    request_assessments(session, sep.touched, sep_doc.id)
    oct_doc, october = import_bytes(session, store, source, PMS_OCT)
    assert len(request_assessments(session, october.touched, oct_doc.id)) == 38
    assert len(_assess_jobs(session)) == 76


def test_a_changed_state_value_also_queues_the_national_situation(
    session: Session, store: S3Store, source: Source, places: dict[str, int]
) -> None:
    doc, _ = import_bytes(session, store, source, PMS_OCT)
    request_assessments(session, {("pms_litre", places["NG-LA"])}, doc.id + 1000)
    national = _situation(session, "pms_litre", "NG")
    lagos = _situation(session, "pms_litre", "NG-LA")
    keys = set(_assess_jobs(session))
    assert keys == {
        f"assess_situation:{national.id}:{doc.id + 1000}",
        f"assess_situation:{lagos.id}:{doc.id + 1000}",
    }


def _place(session: Session, code: str) -> int:
    return session.scalars(select(Place.id).where(Place.code == code)).one()
