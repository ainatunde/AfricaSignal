"""Read NBS price-watch Excel files (spec B6.3). No database and no network.

Written against the real files saved in ``tests/fixtures/nbs`` (see ``manifest.json``), not against
an assumed layout. What those files look like:

* Every table has three value columns: the same month a year earlier, the previous month and the
  reference month, in that order. The header names the months either as dates ("2024-10-14", the
  day is meaningless) or as text ("Average of Oct-24").
* Petrol has one state table. Diesel, kerosene and cooking gas also list the six geopolitical
  zones between the states, and end with a NATIONAL or Average row. Kerosene (litre, gallon) and
  cooking gas (5kg, 12.5kg) put two blocks side by side; the block name is in the row above the
  header, over the block's name column.
* The food workbook has one row per item and only a national average: there is no state grid.
* Rows below the national row (highest and lowest states) and the columns to the right of the
  value columns (MoM, YoY) are ignored. MoM and YoY are recomputed from values in code.
* Names are as NBS wrote them: "Abuja" for the FCT, "Nassarawa" and "Nasarawa" both occur.
* Row labels can be wrong. In the October 2024 cooking-gas file the 12.5kg block labels the Kebbi
  row "Taraba". Both blocks share the same rows, so the name in the first block is used and the
  disagreement is reported as a warning.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, Literal

import openpyxl

from africasignal.catalog import NbsPublication

ZONES = frozenset(
    {"north central", "north east", "north west", "south east", "south south", "south west"}
)
NATIONAL_LABELS = frozenset({"national", "average", "grand total"})
_MONTHS = {
    m: i
    for i, m in enumerate(
        ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1
    )
}
HEADER_SEARCH_ROWS = 40
TWO_PLACES = Decimal("0.01")

Column = Literal["year_ago", "previous", "reference"]
COLUMNS: tuple[Column, Column, Column] = ("year_ago", "previous", "reference")


class NbsParseError(ValueError):
    """The workbook does not have the layout the parser was written for."""


@dataclass(frozen=True)
class TableRow:
    name: str  # as written by NBS: a state, an item label, or the national row
    values: dict[Column, Decimal | None]


@dataclass
class ParsedTable:
    sheet: str
    block: str | None  # "5KG", "12.5KG", "LITRE", "GALLON"; None when the table has no blocks
    months: dict[Column, date]  # first day of each column's month
    rows: list[TableRow]  # states (or food items), zones and the national row excluded
    national: TableRow | None  # the NATIONAL or Average row, when the table has one
    warnings: list[str] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    publication: str
    reference_month: date
    tables: list[ParsedTable]

    @property
    def warnings(self) -> list[str]:
        return [w for t in self.tables for w in t.warnings]

    def table(self, block: str | None) -> ParsedTable | None:
        wanted = block.upper() if block else None
        return next((t for t in self.tables if (t.block or None) == wanted), None)


def _month_of(cell: Any) -> date | None:
    """The month a header cell names, or None if it is not a month header."""
    if isinstance(cell, datetime):
        return date(cell.year, cell.month, 1)
    if isinstance(cell, date):
        return date(cell.year, cell.month, 1)
    if isinstance(cell, str) and cell.strip().lower().startswith("average of"):
        token = cell.strip().split()[-1]  # "Oct-24"
        mon, _, yy = token.partition("-")
        if mon[:3].lower() in _MONTHS and yy.isdigit() and len(yy) == 2:
            return date(2000 + int(yy), _MONTHS[mon[:3].lower()], 1)
    return None


def _add_months(d: date, n: int) -> date:
    index = d.year * 12 + (d.month - 1) + n
    return date(index // 12, index % 12 + 1, 1)


def _decimal(cell: Any) -> Decimal | None:
    if isinstance(cell, bool) or not isinstance(cell, int | float | Decimal):
        return None
    return Decimal(repr(cell) if isinstance(cell, float) else str(cell)).quantize(
        TWO_PLACES, rounding=ROUND_HALF_UP
    )


def _text(cell: Any) -> str | None:
    if cell is None:
        return None
    value = str(cell).strip()
    return value or None


def _find_header(rows: list[tuple[Any, ...]]) -> tuple[int, list[int]] | None:
    """The first row holding runs of three month cells, and the column where each run starts."""
    for r, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        starts = [
            c
            for c in range(len(row) - 2)
            if all(_month_of(row[c + k]) is not None for k in range(3))
        ]
        if starts:
            return r, starts
    return None


def _cell(rows: list[tuple[Any, ...]], r: int, c: int) -> Any:
    return rows[r][c] if 0 <= r < len(rows) and 0 <= c < len(rows[r]) else None


def _block_name(rows: list[tuple[Any, ...]], header: int, start: int) -> str | None:
    """The block a value run belongs to ("5KG"): text in the row above the header, over the
    block's name column (2024 files) or over its first value column (2026 files). A row of month
    headers above is another header, not a block name."""
    if header == 0 or _month_of(_cell(rows, header - 1, start)) is not None:
        return None
    for col in (start - 1, start):
        name = _text(_cell(rows, header - 1, col))
        if name is not None:
            return name.upper()
    return None


def _read_table(
    sheet: str, rows: list[tuple[Any, ...]], header: int, start: int, name_col: int
) -> ParsedTable:
    year_ago, previous, reference = (_month_of(_cell(rows, header, start + k)) for k in range(3))
    if year_ago is None or previous is None or reference is None:
        raise NbsParseError(f"{sheet}: header at row {header + 1} has no month columns")
    if previous != _add_months(reference, -1) or year_ago != _add_months(reference, -12):
        raise NbsParseError(
            f"{sheet}: value columns are {year_ago:%b %Y}, {previous:%b %Y}, {reference:%b %Y}; "
            "expected the same month last year, the previous month and the reference month"
        )

    block = _block_name(rows, header, start)
    table = ParsedTable(
        sheet=sheet,
        block=block,
        months={"year_ago": year_ago, "previous": previous, "reference": reference},
        rows=[],
        national=None,
    )
    seen: set[str] = set()
    for r in range(header + 1, len(rows)):
        values: dict[Column, Decimal | None] = {
            col: _decimal(_cell(rows, r, start + k)) for k, col in enumerate(COLUMNS)
        }
        name = _text(_cell(rows, r, name_col))
        if name is None:
            if table.rows and all(v is None for v in values.values()):
                break  # the blank row that ends the table
            continue
        if start - 1 != name_col:
            other = _text(_cell(rows, r, start - 1))
            if other is not None and other.lower() != name.lower():
                table.warnings.append(
                    f"{sheet} row {r + 1}: block {block} labels this row {other!r} but the first "
                    f"block labels it {name!r}; used {name!r}"
                )
        key = name.lower()
        if key in ZONES:
            continue
        if key in NATIONAL_LABELS:
            table.national = TableRow(name, values)
            break  # the highest and lowest states are listed below; they are not data
        if key in seen:
            raise NbsParseError(f"{sheet} row {r + 1}: {name!r} appears twice in one table")
        seen.add(key)
        if all(v is None for v in values.values()):
            continue  # a stray label, not data
        table.rows.append(TableRow(name, values))
    if not table.rows:
        raise NbsParseError(f"{sheet}: no data rows under the header at row {header + 1}")
    return table


def _wanted(
    publication: NbsPublication, rows: list[tuple[Any, ...]], header: int, start: int
) -> bool:
    """Whether the table whose value columns start at ``start`` belongs to the publication."""
    label = (_text(_cell(rows, header, start - 1)) or "").lower()
    has_block = _block_name(rows, header, start) is not None
    match publication.layout:
        case "state_table":  # petrol: the header cell above the names says "State(s)"
            return label in ("state", "states")
        case "zoned_table":  # diesel: no header label (2024) or "STATES" (2026), no block name
            return label in ("", "states") and not has_block
        case "two_blocks":  # kerosene, cooking gas: a block name sits above the header
            return has_block
        case "national_items":  # food: the header cell says "Items Label"
            return label.startswith("item")
    return False


# An xlsx file is a ZIP of XML parts, and a few kilobytes of it can expand to gigabytes. The real
# NBS workbooks are well under a megabyte unpacked; anything far past that is refused unread.
MAX_UNPACKED_BYTES = 50_000_000
MAX_ZIP_MEMBERS = 500
# A sheet declares its own extent, and one cell at XFD1048576 makes ``iter_rows`` walk seventeen
# billion empty cells. The real tables are under 100 rows by 20 columns.
MAX_SHEETS = 50
MAX_SHEET_ROWS = 5_000
MAX_SHEET_COLUMNS = 200


def check_zip_size(content: bytes) -> None:
    """Raise ``NbsParseError`` when the ZIP in ``content`` would unpack to an unreasonable size
    (by the sizes its directory declares; Python reads no more than a member declares). Content
    that is not a ZIP is left for the caller to reject."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
    except zipfile.BadZipFile:
        return
    if len(members) > MAX_ZIP_MEMBERS or sum(m.file_size for m in members) > MAX_UNPACKED_BYTES:
        raise NbsParseError("the file unpacks to an unreasonable size and was not opened")


def parse_workbook(content: bytes, publication: NbsPublication) -> ParsedWorkbook:
    """Read every price table in the workbook that belongs to ``publication``.

    A workbook can carry sheets of other publications (the October 2024 petrol file also holds
    the diesel sheet), so tables are chosen by layout, not by sheet order.
    """
    check_zip_size(content)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:
        raise NbsParseError(f"not a readable Excel workbook: {exc}") from exc

    if len(wb.worksheets) > MAX_SHEETS:
        raise NbsParseError("the workbook has too many sheets and was not read")
    for ws in wb.worksheets:
        if (ws.max_row or 0) > MAX_SHEET_ROWS or (ws.max_column or 0) > MAX_SHEET_COLUMNS:
            raise NbsParseError(f"sheet {ws.title!r} is far larger than a price table; not read")

    tables: list[ParsedTable] = []
    for ws in wb.worksheets:
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        found = _find_header(rows)
        if found is None:
            continue
        header, starts = found
        tables += [
            _read_table(ws.title, rows, header, start, starts[0] - 1)
            for start in starts
            if _wanted(publication, rows, header, start)
        ]
        if tables and publication.layout != "two_blocks":
            break  # one table per workbook; the first sheet that matches wins
    if not tables:
        raise NbsParseError(
            f"no {publication.code} price table found in sheets "
            f"{[ws.title for ws in wb.worksheets]}"
        )

    references = {t.months["reference"] for t in tables}
    if len(references) != 1:
        raise NbsParseError(f"tables disagree on the reference month: {sorted(references)}")
    return ParsedWorkbook(publication.code, references.pop(), tables)
